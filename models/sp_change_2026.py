# -*- coding: utf-8 -*-
"""SP 2026 智能体自动填写 - 业务核心。

解析驱动文件 (gg/战略规划SP变更需求-*.xlsx) -> 过滤"智能体自动..."行 ->
对每行调 MCP 拉 2025 实际数据 (few-shot) -> LLM + 联网生成 2026 单元格数据。
"""
from __future__ import annotations
import os
import re
from dataclasses import dataclass, asdict
from openpyxl import load_workbook
from models import idste


AUTO_FILL_MARKERS = {"智能体自动搜索填写", "智能体自动分析总结"}


@dataclass
class RequirementRow:
    """驱动文件里一行 (BLM 战略规划研究内容追踪表)。"""
    row: int
    blm_module: str
    stage_step: str
    research_content: str
    output_template: str
    data_source: str
    process_hint: str
    idste_module: str
    note: str

    def to_dict(self) -> dict:
        return asdict(self)


def parse_requirement_xlsx(path: str) -> list[RequirementRow]:
    """读驱动文件 Excel, 过滤 E in AUTO_FILL_MARKERS 的行。"""
    if not os.path.exists(path):
        raise FileNotFoundError(f"驱动文件不存在: {path}")
    if not path.lower().endswith((".xlsx", ".xlsm")):
        raise ValueError(f"驱动文件必须是 .xlsx: {path}")

    wb = load_workbook(path, data_only=True)
    ws = wb[wb.sheetnames[0]]

    rows: list[RequirementRow] = []
    for r_idx in range(3, ws.max_row + 1):
        a = ws.cell(row=r_idx, column=1).value
        e = ws.cell(row=r_idx, column=5).value
        if e is None:
            continue
        e_str = str(e).strip()
        if e_str not in AUTO_FILL_MARKERS:
            continue
        rows.append(RequirementRow(
            row=r_idx,
            blm_module=str(a or "").strip(),
            stage_step=str(ws.cell(row=r_idx, column=2).value or "").strip(),
            research_content=str(ws.cell(row=r_idx, column=3).value or "").strip(),
            output_template=str(ws.cell(row=r_idx, column=4).value or "").strip(),
            data_source=e_str,
            process_hint=str(ws.cell(row=r_idx, column=6).value or "").strip(),
            idste_module=str(ws.cell(row=r_idx, column=7).value or "").strip(),
            note=str(ws.cell(row=r_idx, column=8).value or "").strip(),
        ))
    return rows


def _normalize_template_name(name: str) -> str:
    """去"表"字 + 去空白, 方便与 menu 里 table.name 模糊匹配。"""
    return re.sub(r"\s+", "", name).rstrip("表")


def resolve_table_key(template_name: str, menu, keyword: str | None = None) -> str:
    """把 D 列里的"输出表格/图形模版"名解析为 MCP table_key。

    匹配策略 (按优先级):
    1. 评分匹配: 对每个候选, 找其 name 包含的**最长** CJK chunk (来自 template_name),
       候选得分 = 最长 chunk 长度。取所有候选中最高分。这避免短通用 chunk 如 "分析"
       误匹配 (例: "2.1 宏观环境分析(PESTEL分析)表" 不会误匹配到 "行业趋势分析" -- 因为
       "宏观环境分析"(6字) 得分高于 "分析"(2字))。
    2. 评分无命中时, 用 CJK 字符集重叠打分: 候选得分 = 候选 name 中出现的
       normalized CJK 字符数 (去重)。处理 MCP name 用"/"分隔的场景 (如"看行业/趋势"
       对应 template "2.2 行业趋势分析": 行/业/趋/势 4 字命中)。
    3. 仍无命中 -> ValueError。
    """
    candidates = idste._find_table_keys(menu)
    if not candidates:
        raise ValueError(
            f"sp_data_menu 返回里未发现任何 table_key, 无法解析 {template_name!r}"
        )

    normalized = _normalize_template_name(template_name)
    if keyword is None:
        keyword = "".join(re.findall(r"[一-鿿]+", normalized))

    # --- 1. 评分匹配 (longest CJK chunk wins; 防短通用 chunk 误匹配) ---
    # chunks 去重 + 按长度降序, 保证每个候选的"最长匹配"是第一个找到的
    chunks = sorted(set(re.findall(r"[一-鿿]{2,}", normalized)), key=len, reverse=True)
    best_key: str | None = None
    best_len = 0
    for c in candidates:
        cname = re.sub(r"\s+", "", c.get("name") or "")
        for chunk in chunks:  # 已降序, 第一个匹配即该候选的最长匹配
            if chunk in cname:
                if len(chunk) > best_len:
                    best_len = len(chunk)
                    best_key = c["table_key"]
                break  # 该候选最长匹配已找, 看下一个候选
    if best_key is not None:
        return best_key

    # --- 2. fallback: CJK 字符集重叠打分 (处理 name 用"/"等分隔的场景) ---
    norm_cjk = set(re.findall(r"[一-鿿]", normalized))
    if norm_cjk:
        # 阈值: 至少命中 normalized 一半的 CJK 字符 (防"分析"等公共字误匹配)
        threshold = max(2, len(norm_cjk) // 2)
        best_key = None
        best_score = 0
        for c in candidates:
            cname = c.get("name") or ""
            cand_cjk = set(re.findall(r"[一-鿿]", cname))
            score = len(norm_cjk & cand_cjk)  # 交集大小
            if score >= threshold and score > best_score:
                best_score = score
                best_key = c["table_key"]
        if best_key is not None:
            return best_key

    raise ValueError(
        f"主模板 sheet {template_name!r} (normalized={normalized!r}, "
        f"keyword={keyword!r}) 未在 MCP 菜单中找到。\n"
        f"MCP 菜单候选: {[c.get('name') for c in candidates[:8]]}"
    )


if __name__ == "__main__":
    import sys

    DRIVE_FILE = r"D:\Desktop\gg\战略规划SP变更需求-6.5 版（产品需求沟通确认 2.0）.xlsx"

    print("=" * 60)
    print("== [1] parse_requirement_xlsx 离线单元测试 ==")
    rows = parse_requirement_xlsx(DRIVE_FILE)
    print(f"驱动文件: {DRIVE_FILE}")
    print(f"发现 {len(rows)} 行「智能体自动...」:")
    for r in rows:
        print(f"  row={r.row} D={r.output_template!r} E={r.data_source!r}")

    print()
    print("=" * 60)
    print("== [2] resolve_table_key 离线 mock 测试 (防短 chunk 误匹配) ==")
    # mock menu 同时含 "宏观环境分析" 和 "行业趋势分析" 候选。
    # 旧算法 (任意 chunk 命中即返回) 会因 chunk "分析" 通用, 把
    # "2.1 宏观环境分析(PESTEL分析)表" 误匹配到 "行业趋势分析" (若其排在前面)。
    # 新算法 (longest chunk 评分) 应选 "宏观环境分析" (6 字) 而非 "分析" (2 字)。
    mock_menu = {
        "tables": [
            {"table_key": "tk_industry", "name": "2.2 行业趋势分析"},  # 故意放第一
            {"table_key": "tk_pestel", "name": "2.1 宏观环境分析"},
            {"table_key": "tk_market", "name": "2.3 市场容量分析"},
        ]
    }
    # case 1: PESTEL 表不应误匹配到行业趋势 (核心防回归测试)
    try:
        tk = resolve_table_key("2.1 宏观环境分析(PESTEL分析)表", mock_menu)
        assert tk == "tk_pestel", f"期望 tk_pestel, 实际 {tk}"
        print(f"  [PASS] '2.1 宏观环境分析(PESTEL分析)表' -> {tk} (未误匹配到行业趋势)")
    except AssertionError as e:
        print(f"  [FAIL] {e}")
    except ValueError as e:
        print(f"  [FAIL] 未找到: {e}")

    # case 2: 正常匹配仍工作
    try:
        tk = resolve_table_key("2.2 行业趋势分析表", mock_menu)
        assert tk == "tk_industry", f"期望 tk_industry, 实际 {tk}"
        print(f"  [PASS] '2.2 行业趋势分析表' -> {tk}")
    except (AssertionError, ValueError) as e:
        print(f"  [FAIL] {e}")

    # case 3: 无匹配应抛 ValueError
    try:
        resolve_table_key("9.9 不存在的表", mock_menu)
        print(f"  [FAIL] 应抛 ValueError 但没抛")
    except ValueError:
        print(f"  [PASS] 无匹配正确抛 ValueError")

    print()
    print("=" * 60)
    print("== [3] resolve_table_key 在线测试 (需 VPN/iDSTE) ==")
    if len(sys.argv) > 1 and sys.argv[1] == "--online":
        menu = idste.sp_data_menu("c")
        for r in rows[:3]:
            try:
                tk = resolve_table_key(r.output_template, menu)
                print(f"  {r.output_template} -> {tk} OK")
            except ValueError as e:
                print(f"  {r.output_template} -> FAIL: {e}")
    else:
        print("  (跳过; 跑 `python -m models.sp_change_2026 --online` 启用)")
