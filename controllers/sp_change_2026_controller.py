# -*- coding: utf-8 -*-
"""SP 2026 智能体自动填写 - 端到端编排控制器。

驱动文件 (gg/战略规划SP变更需求-*.xlsx) -> 过滤"智能体自动..."行 ->
对每行: MCP 拉 2025 实际数据 (few-shot) + 读主模板 A 列维度标签 ->
build_sp_2026_data_prompt -> LLM agent loop (web_search 工具) ->
解析 JSON cells -> in-place 写入主模板副本。

关键设计:
- sp_dimension / sp_data_menu 在 loop 外**只调一次** (避免 N+1 次 MCP 调用)
- out_xlsx_path 第一次写入时创建 (复制主模板), 后续 sheet 累加 in-place
- 每个 sheet 失败不中断, 累计到 errors[]; 全部失败才 raise RuntimeError
- 所有依赖 (llm_client / tool_map / prompts_module / sp_change_2026_module /
  xlsx_exporter_module) 都参数化, 便于测试与多账号/多模型替换
"""
from __future__ import annotations

import os
import re
import json
from datetime import datetime
from uuid import uuid4

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string

from models import idste
from models import llm
from models import prompts
from models import websearch
from models import sp_change_2026 as sp_change_2026_mod
from models import xlsx_exporter
from core.envutil import env


# ============================================================
# 工具函数
# ============================================================
def _master_template_sheet(output_template: str, available_sheets: list[str] | None = None) -> str:
    """驱动文件 D 列 (如 '2.1 宏观环境分析(PESTEL分析)表') -> 主模板 sheet 名 (如 '2.1 宏观环境分析')。

    Args:
        output_template: D 列原值
        available_sheets: 主模板实际 sheet 名列表; 提供时做模糊查 (容错 D 列尾部的
            "表"/"图"字 + "(...)" 注释), 否则仅去尾部"表"字。

    匹配优先级 (仅在 available_sheets 提供时):
    1) 精确 (strip 后)
    2) 前缀: D 列名以 sheet 名开头 (容错 D 列尾部 "(PESTEL分析)" 等注释)
    3) 共同 CJK chunk (longest wins, 同 resolve_table_key 思路)
    """
    raw = output_template.strip()
    # 去尾部"表"/"图"字
    s = raw
    for suffix in ("表", "图"):
        if s.endswith(suffix):
            s = s[:-len(suffix)].strip()
            break
    if not available_sheets:
        return s
    # 1) 精确 (strip 后)
    for name in available_sheets:
        if name.strip() == s:
            return name
    # 2) 前缀匹配 (D 列名以 sheet 名开头, 容错 D 列尾部的 "(...)" 注释)
    for name in available_sheets:
        n = name.strip()
        if n and s.startswith(n):
            return name
    # 3) 共同 CJK chunk (longest wins)
    chunks = sorted(set(re.findall(r"[一-鿿]{2,}", s)), key=len, reverse=True)
    best: str | None = None
    best_len = 0
    for name in available_sheets:
        n = name.strip()
        for chunk in chunks:
            if chunk in n and len(chunk) > best_len:
                best = name
                best_len = len(chunk)
                break
    return best if best is not None else s


def _read_row_labels_from_template(
    template_path: str,
    sheet_name: str,
    target_col: str = "A",
    start_row: int = 4,
) -> list[str]:
    """从主模板 sheet 的 target_col 从 start_row 起读连续非空 cell 值, 作为维度标签。

    遇第一个空 cell 停止 (或到 max_row)。

    Raises:
        KeyError: sheet_name 不在模板里
    """
    wb = load_workbook(template_path, data_only=True)
    available = {name.strip(): name for name in wb.sheetnames}
    sname_norm = sheet_name.strip()
    if sname_norm not in available:
        raise KeyError(
            f"主模板里没有 sheet {sheet_name!r}。已有: {list(wb.sheetnames)[:5]}..."
        )
    ws = wb[available[sname_norm]]
    col_idx = column_index_from_string(target_col)

    labels: list[str] = []
    r = start_row
    while r <= ws.max_row:
        v = ws.cell(row=r, column=col_idx).value
        if v is None or str(v).strip() == "":
            break
        labels.append(str(v).strip())
        r += 1
    return labels


def _sp_data_to_baseline_text(data) -> str:
    """JSON 序列化 sp_data 返回, 截断到 20000 字 (防止 prompt 超长)。"""
    s = json.dumps(data, ensure_ascii=False, default=str, indent=2)
    if len(s) > 20000:
        s = s[:20000] + "\n...(已截断)..."
    return s


def _extract_json_from_text(text: str) -> dict:
    """容错解析 LLM 最终回答为 JSON dict (沿用 controllers.agent._parse_json 思路)。

    处理: ```json 围栏 / 前后言包裹 / 裸引号修复。失败返回 {}, 绝不抛异常。

    注: agent._parse_json 旧版 rstrip("`") 后比较, 导致闭合 ``` 变成 "" 不被过滤,
    本函数改用 strip() 后直接比较, 修了这一处。
    """
    from controllers.agent import _repair_unescaped_quotes

    t = text.strip()
    # 去 markdown 围栏 (```json ... ``` 或 ``` ... ```)
    if t.startswith("```"):
        lines = t.split("\n")
        # 丢首尾 fence 行 (用 strip 后比较, 不 rstrip("`") 避免 agent._parse_json 旧 bug)
        kept = [ln for ln in lines if ln.strip().lower() not in ("```json", "```", "``")]
        t = "\n".join(kept).strip()
    # 提取第一个 {...}
    if not t.startswith("{"):
        i = t.find("{")
        j = t.rfind("}")
        if i != -1 and j != -1:
            t = t[i:j + 1]
    try:
        obj = json.loads(t)
    except json.JSONDecodeError:
        # fallback: 修复字符串值内未转义的双引号 (reasoning 模型常见问题)
        try:
            obj = json.loads(_repair_unescaped_quotes(t))
        except Exception:
            return {}
    return obj if isinstance(obj, dict) else {}


def _pick_dim_info(dim_resp) -> str:
    """从 sp_dimension() 返回中提取 dim_info 字符串。优先 'c' (公司维度), 否则第一个。

    容错多种返回结构: str / list[str|dict] / dict{dimensions: [...]}。
    """
    if isinstance(dim_resp, str):
        return dim_resp
    items: list = []
    if isinstance(dim_resp, list):
        items = dim_resp
    elif isinstance(dim_resp, dict):
        for k in ("dimensions", "data", "list", "items"):
            if isinstance(dim_resp.get(k), list):
                items = dim_resp[k]
                break
    dims: list[str] = []
    for it in items:
        if isinstance(it, str):
            dims.append(it)
        elif isinstance(it, dict):
            v = it.get("dim_info") or it.get("id") or it.get("code")
            if v:
                dims.append(str(v))
    if "c" in dims:
        return "c"
    if dims:
        return dims[0]
    return "c"


def _make_counting_wrapper(name, fn, stats):
    """包一层工具调用计数 (沿用 controllers.agent.generate_xlsx_data 模式)。"""
    def wrapper(**kwargs):
        stats["__all__"].setdefault(name, 0)
        stats["__all__"][name] += 1
        return fn(**kwargs)
    wrapper.__name__ = f"counted_{name}"
    return wrapper


# ============================================================
# 主入口: 端到端编排
# ============================================================
def run_sp_change_2026(
    gg_xlsx_path: str,
    template_xlsx_path: str | None = None,
    year: int = 2026,
    target_sheet_names: list[str] | None = None,
    output_dir: str = "output",
    llm_client=None,
    tool_map=None,
    prompts_module=None,
    sp_change_2026_module=None,
    xlsx_exporter_module=None,
    enable_web_search: bool = True,
    verbose: bool = True,
) -> dict:
    """SP 2026 端到端编排: 驱动文件 -> MCP 2025 基线 -> LLM + web_search -> 写主模板副本。

    Args:
        gg_xlsx_path: 驱动文件 (gg/战略规划SP变更需求-*.xlsx) 路径
        template_xlsx_path: 主模板 xlsx 路径; None 时从 .env 读 TEMPLATE_XLSX_PATH
        year: 目标年份 (默认 2026)
        target_sheet_names: 只填这些 sheet (D 列 output_template 名, 含/不含"表"均可);
            None = 填驱动文件里所有"智能体自动..."行
        output_dir: 输出目录
        llm_client / tool_map / prompts_module / sp_change_2026_module / xlsx_exporter_module:
            依赖注入 (None 时用默认模块)
        verbose: 打印过程

    Returns:
        dict: {
            "intent_meta": {gg_xlsx_path, year, target_sheet_names, n_requirements, dim_info},
            "output_xlsx_path": str | None,  # None 表示全部失败
            "tool_call_stats": {"__all__": {tool_name: count}},
            "filled_sheets": [sheet_name, ...],
            "errors": [{"sheet", "output_template", "error"}, ...],
        }

    Raises:
        RuntimeError: 所有 sheet 都失败 (filled_sheets 为空)
    """
    # ---- DI 默认值 ----
    llm_client = llm_client or llm
    prompts_module = prompts_module or prompts
    sp_change_2026_module = sp_change_2026_module or sp_change_2026_mod
    xlsx_exporter_module = xlsx_exporter_module or xlsx_exporter
    template_xlsx_path = template_xlsx_path or env("TEMPLATE_XLSX_PATH")
    tool_map = tool_map if tool_map is not None else dict(idste.TOOL_MAP)

    # ---- 1. 解析驱动文件 ----
    if verbose:
        print(f"[sp-2026] 解析驱动文件: {gg_xlsx_path}")
    requirements = sp_change_2026_module.parse_requirement_xlsx(gg_xlsx_path)

    # ---- 2. 按 target_sheet_names 过滤 (容错"表"字/空格) ----
    if target_sheet_names:
        target_set = {s.strip().rstrip("表") for s in target_sheet_names}
        requirements = [
            r for r in requirements
            if r.output_template.strip().rstrip("表") in target_set
        ]
    if not requirements:
        raise ValueError(
            f"驱动文件 {gg_xlsx_path} 过滤后无任何 '智能体自动...' 行"
            + (f" (target={target_sheet_names})" if target_sheet_names else "")
        )
    if verbose:
        print(f"[sp-2026] 待填 {len(requirements)} 个 sheet: "
              f"{[r.output_template for r in requirements]}")

    # ---- 3. loop 外一次: sp_dimension + sp_data_menu (避免 N+1 MCP) ----
    if verbose:
        print(f"[sp-2026] 拉 dimension + menu (loop 外一次)...")
    try:
        dim_resp = idste.sp_dimension()
        dim = _pick_dim_info(dim_resp)
    except Exception as e:
        if verbose:
            print(f"[sp-2026] sp_dimension 失败, fallback dim='c': {e}")
        dim = "c"
    menu = idste.sp_data_menu(dim)
    if verbose:
        print(f"[sp-2026] dim={dim!r}, menu 候选 {len(idste._find_table_keys(menu))} 个表")

    # ---- 4. 构造 tool_map (idste + websearch, 全计数) + tools_schema (仅 web_search) ----
    # tool_map 传 idste + websearch (按 brief); schema 只暴露 web_search
    # (LLM 不需要再调 sp_data/sp_dimension -- 2025 基线已 loop 外拉好)
    tool_call_stats: dict = {"__all__": {}}
    full_tool_map: dict = {}
    for name, fn in dict(tool_map).items():
        full_tool_map[name] = _make_counting_wrapper(name, fn, tool_call_stats)
    if enable_web_search:
        for name, fn in websearch.TOOL_MAP.items():
            full_tool_map[name] = _make_counting_wrapper(name, fn, tool_call_stats)
        tools_schema = [websearch.WEB_SEARCH_TOOL_SCHEMA["function"]]
    else:
        tools_schema = []  # 不暴露 web_search; LLM 直接基于 few-shot 生成

    # ---- 4.5 加载主模板 sheet 名一次 (供 mt_sheet 模糊查) ----
    _tpl_wb = load_workbook(template_xlsx_path, read_only=True)
    _tpl_sheet_names = list(_tpl_wb.sheetnames)
    _tpl_wb.close()

    # ---- 5. 输出路径 (第一次写入时由 fill_cells 复制模板创建) ----
    os.makedirs(output_dir, exist_ok=True)
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    out_xlsx_path = os.path.join(
        output_dir, f"{ts}_sp_2026_{uuid4().hex[:8]}_data.xlsx"
    )

    filled_sheets: list[str] = []
    errors: list[dict] = []
    intent_meta = {
        "gg_xlsx_path": gg_xlsx_path,
        "year": year,
        "target_sheet_names": target_sheet_names,
        "n_requirements": len(requirements),
        "dim_info": dim,
    }

    # ---- 6. 逐 sheet 编排 (单 sheet 失败不中断) ----
    for req in requirements:
        mt_sheet = _master_template_sheet(req.output_template, _tpl_sheet_names)
        try:
            # 6.1 resolve_table_key -> MCP table_key
            tk = sp_change_2026_module.resolve_table_key(req.output_template, menu)

            # 6.2 拉 2025 基线 (few-shot)
            baseline_2025 = idste.sp_data(dim, tk, str(year - 1))
            baseline_text = _sp_data_to_baseline_text(baseline_2025)

            # 6.3 读主模板 A 列维度标签
            # 结果列/起始行参数化 (label_col=A 读维度; result_col=B 写结果)
            # 沿用同一对变量到 prompt builder + 缺 cell 回填, 保证扩展非 B 列 sheet 时一致
            label_col = "A"
            result_col = "B"
            result_start_row = 4
            row_labels = _read_row_labels_from_template(
                template_xlsx_path, mt_sheet, target_col=label_col, start_row=result_start_row,
            )
            if not row_labels:
                raise ValueError(
                    f"主模板 sheet {mt_sheet!r} {label_col} 列从 row {result_start_row} 起无维度标签"
                )

            # 6.4 build prompt
            system_prompt = prompts_module.build_sp_2026_data_prompt(
                template_name=mt_sheet,
                row_labels=row_labels,
                baseline_2025_json=baseline_text,
                process_hint=req.process_hint,
                year=year,
                target_col=result_col,
                start_row=result_start_row,
            )
            user_msg = (
                f"按上面规则生成 {year} 年【{mt_sheet}】的 "
                f"{len(row_labels)} 个维度内容。"
                + ("" if enable_web_search
                   else "（注: web_search 不可用, 直接基于 2025 few-shot + 你的行业知识生成 2026 预测, 不要反复尝试联网）")
            )
            if verbose:
                print(f"[sp-2026] -> {mt_sheet}: tk={tk}, dims={len(row_labels)}, "
                      f"baseline={len(baseline_text)}字, prompt={len(system_prompt)}字, "
                      f"web_search={'on' if enable_web_search else 'off'}")

            # 6.5 LLM agent loop (web_search 工具; 关时 LLM 直接出最终答案)
            final_text = llm_client.run_agent_loop(
                system_prompt, user_msg, full_tool_map, tools_schema,
                max_turns=12 if enable_web_search else 2, verbose=verbose,
                max_tokens=2000, final_max_tokens=12288,
                enable_thinking=False, final_enable_thinking=True,
            )

            # 6.6 解析 cells (容错: 缺 cell 填空字符串)
            cells = _extract_json_from_text(final_text)
            if not cells:
                raise ValueError(
                    f"LLM 最终回答解析为空 dict, final_text 长度={len(final_text)}"
                )
            for i in range(len(row_labels)):
                coord = f"{result_col}{result_start_row + i}"
                if coord not in cells:
                    cells[coord] = ""
                    if verbose:
                        print(f"[sp-2026]    警告: {coord} 缺失, 填空字符串")

            # 6.7 in-place 写入 (首次复制模板, 后续 load+save)
            if not os.path.exists(out_xlsx_path):
                # 第一次: 用 fill_cells 复制模板 + 写入 (走 DI'd xlsx_exporter_module)
                xlsx_exporter_module.fill_cells(
                    template_xlsx_path, mt_sheet, cells,
                    mode="overwrite", out_path=out_xlsx_path,
                )
            else:
                # 后续: load + 直接写 cell + save (不动其他 sheet)
                wb = load_workbook(out_xlsx_path)
                available = {name.strip(): name for name in wb.sheetnames}
                sname_norm = mt_sheet.strip()
                if sname_norm not in available:
                    raise KeyError(
                        f"主模板里没有 sheet {mt_sheet!r}。已有: {list(wb.sheetnames)[:5]}..."
                    )
                ws = wb[available[sname_norm]]
                for coord, val in cells.items():
                    ws[coord] = val
                wb.save(out_xlsx_path)

            filled_sheets.append(mt_sheet)
            if verbose:
                print(f"[sp-2026] OK: {mt_sheet} 写入 {len(cells)} cells -> {out_xlsx_path}")

        except Exception as e:
            err = {
                "sheet": mt_sheet,
                "output_template": req.output_template,
                "error": f"{type(e).__name__}: {e}",
            }
            errors.append(err)
            if verbose:
                print(f"[sp-2026] FAIL: {mt_sheet} - {err['error']}")

    # ---- 7. 全部失败 -> raise ----
    if not filled_sheets and errors:
        raise RuntimeError(
            f"所有 {len(requirements)} 个 sheet 都失败: "
            + "; ".join(f"{e['sheet']}({e['error']})" for e in errors)
        )

    return {
        "intent_meta": intent_meta,
        "output_xlsx_path": out_xlsx_path if filled_sheets else None,
        "tool_call_stats": tool_call_stats,
        "filled_sheets": filled_sheets,
        "errors": errors,
    }


# ============================================================
# __main__: 2.2 试跑
# ============================================================
def _test_master_template_sheet_chunk_fallback() -> None:
    """回归测试: 长 chunk 不匹配时, 应 fallback 到短 chunk (验证 break 位置).

    构造 output_template='2.7 雷达图分析(雷达图)' -> chunks=['雷达图分析'(5), '雷达图'(3)]
    (按长度降序). candidate='雷达图数据' 不含长 chunk '雷达图分析', 但含短 chunk '雷达图'.
    旧 bug (break 在 if 外): 只试最长 chunk, 不匹配即跳出 -> 该候选被跳过, 返回原模板名.
    修复后 (break 在 if 内): 长 chunk 不匹配则继续试短 chunk -> 命中 '雷达图', 返回候选.
    """
    got = _master_template_sheet(
        "2.7 雷达图分析(雷达图)",
        available_sheets=["雷达图数据"],
    )
    assert got == "雷达图数据", f"期望 '雷达图数据', 实际 {got!r}"
    print("[test] _master_template_sheet chunk fallback: PASS")


if __name__ == "__main__":
    import sys

    target = sys.argv[1] if len(sys.argv) > 1 else "2.2 行业趋势分析表"
    if target == "test":
        _test_master_template_sheet_chunk_fallback()
        sys.exit(0)
    # --no-web 显式关 web_search (网络封锁时用)
    enable_web_search = "--no-web" not in sys.argv
    gg_path = env("GG_XLSX_PATH") or \
        r"D:\Desktop\gg\战略规划SP变更需求-6.5 版（产品需求沟通确认 2.0）.xlsx"

    print("=" * 60)
    print(f"SP 2026 自动填写: {target}  (web_search={'on' if enable_web_search else 'off'})")
    print(f"驱动文件: {gg_path}")
    print("=" * 60)

    result = run_sp_change_2026(
        gg_xlsx_path=gg_path,
        target_sheet_names=[target],
        enable_web_search=enable_web_search,
        verbose=True,
    )

    print("\n" + "=" * 60)
    print("OK 完成:")
    print(f"  out   : {result['output_xlsx_path']}")
    print(f"  rows  : {len(result['filled_sheets'])}")
    print(f"  errors: {len(result['errors'])}")
    if result["errors"]:
        for e in result["errors"]:
            print(f"  - {e['sheet']}: {e['error']}")
    print(f"  tools : {result['tool_call_stats']['__all__']}")
